"""Workbook and CSV normalization into a sparse, cell-addressed grid.

Rows keep their 1-based spreadsheet row number and every cell keeps its A1
reference, so a value extracted downstream can always be traced back to the
exact cell it came from.
"""

from __future__ import annotations

import csv
import io
from datetime import date, datetime, time
from pathlib import PurePath
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

MAX_ROWS_PER_SHEET = 1000
MAX_COLUMNS_PER_SHEET = 60
MAX_CELL_CHARS = 500


def cell_text(value: Any) -> str | None:
    """Render a cell value the way a reviewer would read it."""
    if value is None:
        return None
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else format(value, ".10g")
    if isinstance(value, datetime | date | time):
        return value.isoformat()
    text = str(value).strip()
    return text[:MAX_CELL_CHARS] if text else None


def build_rows(values: list[list[Any]]) -> tuple[list[dict[str, Any]], bool]:
    """Convert a dense row list into sparse rows with A1 references."""
    rows: list[dict[str, Any]] = []
    truncated = len(values) > MAX_ROWS_PER_SHEET

    for row_index, row_values in enumerate(values[:MAX_ROWS_PER_SHEET], start=1):
        truncated = truncated or len(row_values) > MAX_COLUMNS_PER_SHEET
        cells = []
        for column_index, raw in enumerate(row_values[:MAX_COLUMNS_PER_SHEET], start=1):
            text = cell_text(raw)
            if text is None:
                continue
            cells.append({"column": column_index, "ref": f"{get_column_letter(column_index)}{row_index}", "text": text})
        if cells:
            rows.append({"rowNumber": row_index, "cells": cells})

    return rows, truncated


def normalize_xlsx(payload: bytes, file_name: str) -> dict[str, Any]:
    values_book = load_workbook(io.BytesIO(payload), read_only=False, data_only=True)
    formula_book = load_workbook(io.BytesIO(payload), read_only=False, data_only=False)
    try:
        worksheets: list[dict[str, Any]] = []
        structured_tables: list[dict[str, Any]] = []

        for sheet_index, worksheet in enumerate(values_book.worksheets):
            formula_sheet = formula_book.worksheets[sheet_index]
            values = [list(row) for row in worksheet.iter_rows(values_only=True)]
            rows, truncated = build_rows(values)
            uncached = count_uncached_formulas(formula_sheet, worksheet)
            tables = list(getattr(worksheet, "tables", {}).values())

            worksheets.append(
                {
                    "sheetIndex": sheet_index,
                    "name": worksheet.title,
                    "maxRow": worksheet.max_row,
                    "maxColumn": worksheet.max_column,
                    "tableCount": len(tables),
                    "rows": rows,
                    "metadata": {
                        "sheetState": worksheet.sheet_state,
                        "truncated": truncated,
                        "formulaCellsWithoutCachedValues": uncached,
                    },
                }
            )
            structured_tables.extend(
                {
                    "sheetIndex": sheet_index,
                    "tableKey": table.name,
                    "name": getattr(table, "displayName", None) or table.name,
                    "ref": table.ref,
                    "headerRowCount": getattr(table, "headerRowCount", None),
                    "metadata": {},
                }
                for table in tables
            )

        return {
            "kind": "workbook",
            "workbook": {
                "fileName": file_name,
                "sheetCount": len(values_book.sheetnames),
                "metadata": {
                    "source": "xlsx",
                    "creator": values_book.properties.creator,
                    "title": values_book.properties.title,
                },
                "worksheets": worksheets,
                "structuredTables": structured_tables,
            },
        }
    finally:
        values_book.close()
        formula_book.close()


def count_uncached_formulas(formula_sheet: Any, values_sheet: Any) -> int:
    """Formulas saved without a cached result read back as empty in data-only mode."""
    return sum(
        1
        for row in formula_sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=") and values_sheet[cell.coordinate].value is None
    )


def read_csv_rows(text: str, delimiter: str) -> list[list[str]]:
    """Parse CSV rows, keeping free-text notes below the table as one cell each.

    Spreadsheets exported to CSV often end with notes ("NORMALIZED NOI: $597,400")
    whose unquoted commas would otherwise be split into meaningless cells. Once the
    table has ended with a blank line, a row much narrower than the table header is
    kept as the original line of text.
    """
    physical_lines = text.splitlines()
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    values: list[list[str]] = []
    table_width = 0
    table_ended = False

    for row in reader:
        if not any(cell.strip() for cell in row):
            table_ended = table_ended or table_width > 0
            values.append([])
            continue
        table_width = table_width or len(row)
        if table_ended and len(row) * 2 < table_width:
            values.append([physical_lines[reader.line_num - 1].strip()])
        else:
            values.append(row)
    return values


def normalize_csv(text: str, file_name: str) -> dict[str, Any]:
    sample = "\n".join(text.splitlines()[:10])
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=",\t").delimiter
    except csv.Error:
        delimiter = ","

    values = read_csv_rows(text, delimiter)
    rows, truncated = build_rows(values)
    max_column = max((len(row) for row in values), default=0)

    return {
        "kind": "workbook",
        "workbook": {
            "fileName": file_name,
            "sheetCount": 1,
            "metadata": {"source": "csv", "delimiter": delimiter},
            "worksheets": [
                {
                    "sheetIndex": 0,
                    "name": PurePath(file_name).stem[:100] or "Sheet1",
                    "maxRow": len(values),
                    "maxColumn": max_column,
                    "tableCount": 0,
                    "rows": rows,
                    "metadata": {"source": "csv", "truncated": truncated},
                }
            ],
            "structuredTables": [],
        },
    }


def inspect_xlsx(payload: bytes) -> tuple[dict[str, Any], str]:
    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    try:
        names = workbook.sheetnames
        return (
            {
                "workbookDetected": True,
                "sheetCount": len(names),
                "sheetNames": names,
                "workbookProperties": {"creator": workbook.properties.creator, "title": workbook.properties.title},
            },
            f"Workbook with {len(names)} sheet(s): {', '.join(names)}",
        )
    finally:
        workbook.close()
