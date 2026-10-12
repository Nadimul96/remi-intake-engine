from __future__ import annotations

import base64
import hashlib
import io
import mimetypes
import re
import zipfile
from datetime import datetime, timezone
from email import policy
from email.parser import BytesParser
from email.utils import getaddresses, parsedate_to_datetime
from typing import Any, Literal
from urllib.request import urlopen

from fastapi import FastAPI, HTTPException
from openpyxl import load_workbook
from pydantic import BaseModel, Field
from pypdf import PdfReader


app = FastAPI(title="REMI Extractor", version="0.0.3")


DEFERRED_EXTENSIONS = {"csv", "docx", "pptx", "tsv", "xls", "zip"}
SPREADSHEET_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
THREAD_TOKEN_PATTERN = re.compile(r"<[^>]+>")


class InspectArtifactRequest(BaseModel):
    artifact_blob_id: str
    content_type: str | None = None
    file_name: str
    immutable_url: str
    sha256: str
    size_bytes: int = Field(ge=0)


class InspectArtifactResponse(BaseModel):
    inspection: dict[str, Any]
    normalized_metadata: dict[str, Any]
    preview_text: str | None = None
    review_required: bool = False


class NormalizeArtifactRequest(BaseModel):
    artifact_blob_id: str
    content_type: str | None = None
    file_name: str
    immutable_url: str
    parser_target: Literal["pdf", "eml", "xlsx"]


class NormalizeArtifactResponse(BaseModel):
    content_kind: Literal["document", "email", "workbook"]
    normalized_content: dict[str, Any]


def utc_now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def normalize_extension(file_name: str) -> str | None:
    suffix = file_name.rsplit(".", 1)
    if len(suffix) != 2:
        return None
    return suffix[-1].lower()


def fetch_bytes(url: str) -> bytes:
    with urlopen(url, timeout=30) as response:
        return response.read()


def guess_mime(file_name: str, provided_mime: str | None, payload: bytes) -> str | None:
    if provided_mime and provided_mime != "application/octet-stream":
        return provided_mime

    if payload.startswith(b"%PDF-"):
        return "application/pdf"

    guessed, _ = mimetypes.guess_type(file_name)
    return guessed


def classify_artifact(extension: str | None, mime_type: str | None, payload: bytes) -> dict[str, Any]:
    is_zip_container = zipfile.is_zipfile(io.BytesIO(payload))

    if extension == "pdf" or mime_type == "application/pdf":
        return {
            "fileClassification": "pdf",
            "parserTarget": "pdf",
            "supportTier": "supported",
            "isContainer": False,
        }

    if extension == "eml" or mime_type == "message/rfc822":
        return {
            "fileClassification": "email",
            "parserTarget": "eml",
            "supportTier": "supported",
            "isContainer": False,
        }

    if extension == "xlsx" or mime_type == SPREADSHEET_MIME:
        return {
            "fileClassification": "spreadsheet",
            "parserTarget": "xlsx",
            "supportTier": "supported",
            "isContainer": True,
        }

    if extension in DEFERRED_EXTENSIONS:
        return {
            "fileClassification": "archive" if extension == "zip" else "document",
            "parserTarget": "archive" if extension == "zip" else "none",
            "supportTier": "deferred",
            "isContainer": is_zip_container,
        }

    if is_zip_container:
        return {
            "fileClassification": "archive",
            "parserTarget": "archive",
            "supportTier": "deferred",
            "isContainer": True,
        }

    return {
        "fileClassification": "unknown",
        "parserTarget": "none",
        "supportTier": "unsupported",
        "isContainer": is_zip_container,
    }


def truncate_preview(text: str | None, limit: int = 500) -> str | None:
    if not text:
        return None

    cleaned = " ".join(text.split()).strip()
    if not cleaned:
        return None

    return cleaned[:limit]


def inspect_pdf(payload: bytes) -> tuple[dict[str, Any], str | None]:
    reader = PdfReader(io.BytesIO(payload))
    extracted_fragments: list[str] = []

    for page in reader.pages[: min(3, len(reader.pages))]:
        extracted_fragments.append(page.extract_text() or "")

    preview = truncate_preview(" ".join(extracted_fragments))
    native_text_present = bool(preview)

    return (
        {
            "pageCount": len(reader.pages),
            "nativeTextPresent": native_text_present,
            "ocrNeededLater": not native_text_present,
            "previewSource": "native_text" if native_text_present else "none",
        },
        preview,
    )


def inspect_eml(payload: bytes) -> tuple[dict[str, Any], str | None]:
    message = BytesParser(policy=policy.default).parsebytes(payload)
    attachment_count = 0

    for part in message.walk():
        if part.get_content_disposition() == "attachment":
            attachment_count += 1

    preview = extract_email_body_preview(message)

    return (
        {
            "subject": message.get("subject"),
            "from": message.get("from"),
            "to": message.get("to"),
            "cc": message.get("cc"),
            "sentDate": message.get("date"),
            "attachmentCount": attachment_count,
        },
        preview,
    )


def inspect_xlsx(payload: bytes) -> tuple[dict[str, Any], str | None]:
    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    try:
        preview = truncate_preview(f"Workbook with {len(workbook.sheetnames)} sheet(s): {', '.join(workbook.sheetnames)}")

        return (
            {
                "workbookDetected": True,
                "sheetCount": len(workbook.sheetnames),
                "sheetNames": workbook.sheetnames,
                "workbookProperties": {
                    "creator": workbook.properties.creator,
                    "title": workbook.properties.title,
                },
            },
            preview,
        )
    finally:
        workbook.close()


def extract_pdf_normalized_content(payload: bytes) -> dict[str, Any]:
    reader = PdfReader(io.BytesIO(payload))
    title = reader.metadata.title if reader.metadata else None
    page_entries: list[dict[str, Any]] = []
    document_fragments: list[str] = []

    for index, page in enumerate(reader.pages, start=1):
        page_text = truncate_preview(page.extract_text(), limit=1200)
        text_available = bool(page_text)
        if page_text:
            document_fragments.append(page_text)
        page_entries.append(
            {
                "pageNumber": index,
                "textAvailable": text_available,
                "charCount": len(page_text or ""),
                "textPreview": page_text,
                "metadata": {
                    "previewSource": "native_text" if text_available else "none",
                },
            }
        )

    return {
        "document": {
            "title": title,
            "pageCount": len(reader.pages),
            "textAvailable": any(page["textAvailable"] for page in page_entries),
            "textPreview": truncate_preview(" ".join(document_fragments), limit=1200),
            "metadata": {
                "author": reader.metadata.author if reader.metadata else None,
                "producer": reader.metadata.producer if reader.metadata else None,
                "subject": reader.metadata.subject if reader.metadata else None,
            },
            "pages": page_entries,
        },
        "kind": "document",
    }


def extract_message_addresses(header_value: str | None) -> list[str]:
    if not header_value:
        return []
    return [addr for _, addr in getaddresses([header_value]) if addr]


def extract_email_body_preview(message: Any) -> str | None:
    if message.is_multipart():
        for part in message.walk():
            if part.get_content_disposition() == "attachment":
                continue
            if part.get_content_type() == "text/plain":
                return truncate_preview(part.get_content(), limit=1200)
        return None

    return truncate_preview(message.get_content(), limit=1200)


def normalize_thread_key(message: Any) -> str:
    references = message.get("references")
    in_reply_to = message.get("in-reply-to")
    message_id = message.get("message-id")
    subject = truncate_preview(message.get("subject"), limit=120) or "untitled"

    for header in (references, in_reply_to, message_id):
        if not header:
            continue
        tokens = THREAD_TOKEN_PATTERN.findall(header)
        if tokens:
            return tokens[0]
        return header.strip()

    return f"subject:{subject.lower()}"


def normalize_sent_at(header_value: str | None) -> tuple[str | None, str | None]:
    if not header_value:
        return None, None

    try:
        parsed = parsedate_to_datetime(header_value)
        if parsed.tzinfo is None:
            parsed = parsed.replace(tzinfo=timezone.utc)
        return parsed.astimezone(timezone.utc).isoformat(), header_value
    except Exception:
        return None, header_value


def extract_email_normalized_content(payload: bytes) -> dict[str, Any]:
    message = BytesParser(policy=policy.default).parsebytes(payload)
    body_preview = extract_email_body_preview(message)
    sent_at, sent_at_raw = normalize_sent_at(message.get("date"))
    attachments: list[dict[str, Any]] = []
    part_index = 0

    for part in message.walk():
        filename = part.get_filename()
        disposition = part.get_content_disposition()
        if disposition != "attachment" and not filename:
            continue

        attachment_bytes = part.get_payload(decode=True) or b""
        attachments.append(
            {
                "messagePosition": 0,
                "partIndex": part_index,
                "fileName": filename,
                "contentType": part.get_content_type(),
                "sizeBytes": len(attachment_bytes),
                "contentBase64": base64.b64encode(attachment_bytes).decode("ascii"),
                "contentDisposition": disposition,
                "contentId": part.get("content-id"),
                "sha256": hashlib.sha256(attachment_bytes).hexdigest() if attachment_bytes else None,
                "metadata": {
                    "contentTransferEncoding": part.get("content-transfer-encoding"),
                },
            }
        )
        part_index += 1

    return {
        "email": {
            "thread": {
                "threadKey": normalize_thread_key(message),
                "subject": message.get("subject"),
                "messageCount": 1,
                "metadata": {
                    "hasMultipartBody": message.is_multipart(),
                    "references": message.get("references"),
                },
            },
            "messages": [
                {
                    "position": 0,
                    "messageIdHeader": message.get("message-id"),
                    "subject": message.get("subject"),
                    "sender": message.get("from"),
                    "to": extract_message_addresses(message.get("to")),
                    "cc": extract_message_addresses(message.get("cc")),
                    "sentAt": sent_at,
                    "sentAtRaw": sent_at_raw,
                    "bodyPreview": body_preview,
                    "attachmentCount": len(attachments),
                    "metadata": {
                        "replyTo": message.get("reply-to"),
                    },
                }
            ],
            "attachments": attachments,
        },
        "kind": "email",
    }


def extract_xlsx_normalized_content(payload: bytes, file_name: str) -> dict[str, Any]:
    workbook = load_workbook(io.BytesIO(payload), read_only=False, data_only=True)
    try:
        worksheet_entries: list[dict[str, Any]] = []
        structured_tables: list[dict[str, Any]] = []

        for sheet_index, worksheet in enumerate(workbook.worksheets):
            table_map = getattr(worksheet, "tables", {})
            table_values = list(table_map.values()) if table_map else []
            worksheet_entries.append(
                {
                    "sheetIndex": sheet_index,
                    "name": worksheet.title,
                    "maxRow": worksheet.max_row,
                    "maxColumn": worksheet.max_column,
                    "tableCount": len(table_values),
                    "metadata": {
                        "sheetState": worksheet.sheet_state,
                    },
                }
            )

            for table in table_values:
                structured_tables.append(
                    {
                        "sheetIndex": sheet_index,
                        "tableKey": table.name,
                        "name": table.displayName if getattr(table, "displayName", None) else table.name,
                        "ref": table.ref,
                        "headerRowCount": table.headerRowCount if hasattr(table, "headerRowCount") else None,
                        "metadata": {},
                    }
                )

        return {
            "kind": "workbook",
            "workbook": {
                "fileName": file_name,
                "sheetCount": len(workbook.sheetnames),
                "metadata": {
                    "creator": workbook.properties.creator,
                    "title": workbook.properties.title,
                },
                "worksheets": worksheet_entries,
                "structuredTables": structured_tables,
            },
        }
    finally:
        workbook.close()


@app.get("/health")
def health() -> dict[str, str]:
    return {"service": "remi-extractor", "status": "ok"}


@app.post("/inspect", response_model=InspectArtifactResponse)
def inspect_artifact(request: InspectArtifactRequest) -> InspectArtifactResponse:
    try:
        payload = fetch_bytes(request.immutable_url)
    except Exception as error:  # pragma: no cover - transport failure
        raise HTTPException(status_code=502, detail=f"Failed to fetch immutable artifact: {error}") from error

    extension = normalize_extension(request.file_name)
    mime_type = guess_mime(request.file_name, request.content_type, payload)
    classification = classify_artifact(extension, mime_type, payload)

    format_metadata: dict[str, Any] = {}
    preview: str | None = None

    if classification["parserTarget"] == "pdf":
        format_metadata, preview = inspect_pdf(payload)
        if not format_metadata["nativeTextPresent"]:
            classification["supportTier"] = "best_effort"
    elif classification["parserTarget"] == "eml":
        format_metadata, preview = inspect_eml(payload)
    elif classification["parserTarget"] == "xlsx":
        format_metadata, preview = inspect_xlsx(payload)

    inspection_status = "completed"
    if classification["supportTier"] in {"deferred", "unsupported"}:
        inspection_status = "unsupported"

    inspection = {
        "artifactBlobId": request.artifact_blob_id,
        "detectedMimeType": mime_type,
        "fileClassification": classification["fileClassification"],
        "inspectionStatus": inspection_status,
        "inspectedAt": utc_now_iso(),
        "isContainer": classification["isContainer"],
        "nativeTextPresent": format_metadata.get("nativeTextPresent"),
        "normalizedExtension": extension,
        "ocrNeededLater": format_metadata.get("ocrNeededLater"),
        "parserEligible": classification["supportTier"] in {"supported", "best_effort"},
        "parserTarget": classification["parserTarget"],
        "sizeBytes": request.size_bytes,
        "sourceContentType": request.content_type,
        "supportTier": classification["supportTier"],
    }

    normalized_metadata = {
        "artifactBlobId": request.artifact_blob_id,
        "detectedMimeType": mime_type,
        "fileName": request.file_name,
        "formatMetadata": format_metadata,
        "parserTarget": classification["parserTarget"],
        "sha256": request.sha256,
        "supportTier": classification["supportTier"],
    }

    return InspectArtifactResponse(
        inspection=inspection,
        normalized_metadata=normalized_metadata,
        preview_text=preview,
        review_required=False,
    )


@app.post("/normalize", response_model=NormalizeArtifactResponse)
def normalize_artifact(request: NormalizeArtifactRequest) -> NormalizeArtifactResponse:
    try:
        payload = fetch_bytes(request.immutable_url)
    except Exception as error:  # pragma: no cover - transport failure
        raise HTTPException(status_code=502, detail=f"Failed to fetch immutable artifact: {error}") from error

    if request.parser_target == "pdf":
        normalized_content = extract_pdf_normalized_content(payload)
        return NormalizeArtifactResponse(content_kind="document", normalized_content=normalized_content)

    if request.parser_target == "eml":
        normalized_content = extract_email_normalized_content(payload)
        return NormalizeArtifactResponse(content_kind="email", normalized_content=normalized_content)

    if request.parser_target == "xlsx":
        normalized_content = extract_xlsx_normalized_content(payload, request.file_name)
        return NormalizeArtifactResponse(content_kind="workbook", normalized_content=normalized_content)

    raise HTTPException(status_code=400, detail=f"Unsupported parser target: {request.parser_target}")


if __name__ == "__main__":
    import uvicorn

    uvicorn.run(app, host="0.0.0.0", port=4003)
